"""Extract small, traceable CSVs from the official published Excel tables.

Requires openpyxl. No workbook is edited. Run from any directory:
  python scripts/prepare_data.py --source-dir "../tmp/research"
  python scripts/prepare_data.py --download
"""

import argparse
import csv
import hashlib
import io
import json
import re
from pathlib import Path
from urllib.request import Request, urlopen
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SOURCES = {
    "qilt_ses_2025": {
        "file": "SES_2025_National_Report_Tables.xlsx",
        "url": "https://www.qilt.edu.au/docs/default-source/default-document-library/ses_2025_national_report_tables.zip?sfvrsn=1811247e_1",
        "page": "https://www.qilt.edu.au/surveys/student-experience-survey-(ses)",
        "publisher": "Social Research Centre for the Australian Government Department of Education",
        "reference_year": 2025,
        "publication": "September 2026",
    },
    "education_students_2024": {
        "file": "education_2024_all_students.xlsx",
        "url": "https://www.education.gov.au/download/19481/2024-section-2-all-students/41929/document/xlsx",
        "page": "https://www.education.gov.au/higher-education-statistics/resources/2024-section-2-all-students",
        "publisher": "Australian Government Department of Education",
        "reference_year": 2024,
        "publication": "September 2025",
    },
}
METRICS = ["skills", "peer_engagement", "teaching", "support", "resources", "overall"]
LABELS = {
    "skills": "Skills Development",
    "peer_engagement": "Peer Engagement",
    "teaching": "Teaching Quality and Engagement",
    "support": "Student Support and Services",
    "resources": "Learning Resources",
    "overall": "Overall Educational Experience",
}
STATES = [
    "New South Wales", "Victoria", "Queensland", "Western Australia",
    "South Australia", "Tasmania", "Northern Territory", "Australian Capital Territory",
]
NUMBER = re.compile(r"^(\d+(?:\.\d+)?)(?:\s*\((\d+(?:\.\d+)?),\s*(\d+(?:\.\d+)?)\))?$")


def parse_score(raw):
    """Retain missing/suppressed values, never replace them with zero."""
    if raw is None or str(raw).strip() == "":
        return {"positive_pct": None, "ci_low": None, "ci_high": None, "status": "unavailable"}
    match = NUMBER.fullmatch(str(raw).strip())
    if not match:
        if str(raw).strip().lower() in {"np", "n.p.", "n/a", "-", "*", "**"}:
            return {"positive_pct": None, "ci_low": None, "ci_high": None, "status": str(raw)}
        raise ValueError(f"Unrecognised survey value: {raw!r}")
    value, lo, hi = (float(x) if x is not None else None for x in match.groups())
    assert 0 <= value <= 100
    if lo is not None:
        assert 0 <= lo <= value <= hi <= 100
    return {"positive_pct": value, "ci_low": lo, "ci_high": hi, "status": "published"}


def count(raw):
    if isinstance(raw, (int, float)):
        assert raw >= 0 and int(raw) == raw
        return int(raw), "published"
    return None, "unavailable" if raw is None else str(raw).strip()


def clean_name(raw):
    return re.sub(r"\(\d+\.\d+\)", "", str(raw)).strip().rstrip("*").strip()


def name_key(raw):
    value = clean_name(raw).casefold()
    if value.startswith("the "):
        value = value[4:]
    return {"cquniversity": "central queensland university",
            "torrens university australia": "torrens university"}.get(value, value)


def body(sheet):
    """SES tables use row 4 headers followed by a contiguous data block."""
    for row in sheet.iter_rows(min_row=5):
        if row[1].value is None or str(row[1].value).strip() == "":
            break
        yield row


def trace(sheet, cell):
    return {"source_sheet": sheet.title, "source_cell": cell.coordinate}


def write_csv(out, name, rows):
    assert rows, f"Empty output: {name}"
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with (out / name).open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    return {"file": name, "rows": len(rows), "bytes": (out / name).stat().st_size}


def download(source_dir):
    source_dir.mkdir(parents=True, exist_ok=True)
    for source in SOURCES.values():
        destination = source_dir / source["file"]
        if destination.exists():
            continue
        with urlopen(Request(source["url"], headers={"User-Agent": "FIT3179 student data visualisation"}), timeout=60) as response:
            content = response.read()
        if source["url"].split("?")[0].endswith(".zip"):
            with ZipFile(io.BytesIO(content)) as archive:
                content = archive.read(source["file"])
        destination.write_bytes(content)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=ROOT / "data" / "raw")
    parser.add_argument("--download", action="store_true")
    args = parser.parse_args()
    if args.download:
        download(args.source_dir)
    for source in SOURCES.values():
        path = args.source_dir / source["file"]
        if not path.exists():
            parser.error(f"Missing {path}. Supply --source-dir or --download.")
        source["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()

    qilt = load_workbook(args.source_dir / SOURCES["qilt_ses_2025"]["file"], read_only=True, data_only=True)
    education = load_workbook(args.source_dir / SOURCES["education_students_2024"]["file"], read_only=True, data_only=True)
    out = ROOT / "data"
    out.mkdir(exist_ok=True)
    outputs = []
    used_sheets = set()
    universities, benchmarks, areas, groups, belonging, items, leaving = [], [], [], [], [], [], []

    for code, level in [("UG", "Undergraduate"), ("PGC", "Postgraduate coursework")]:
        sheet = qilt[f"FOCUS_{code}_UNI_17-YY_INST_CI"]
        used_sheets.add(sheet.title)
        for row in body(sheet):
            if clean_name(row[1].value) == "All Universities":
                continue
            for column in range(2, 56):
                metric = METRICS[(column - 2) // 9]
                year = 2017 + (column - 2) % 9
                if year != 2025 and metric != "overall":
                    continue
                cell = row[column]
                universities.append({"institution": clean_name(row[1].value), "level": level,
                    "year": year, "metric": metric, **parse_score(cell.value), **trace(sheet, cell)})

        sheet = qilt[f"FOCUS_{code}_ALL_17-YY_HEPTYPE"]
        used_sheets.add(sheet.title)
        for row in body(sheet):
            for column in range(2, 56):
                cell = row[column]
                benchmarks.append({"provider_type": row[1].value, "level": level,
                    "year": 2017 + (column - 2) % 9, "metric": METRICS[(column - 2) // 9],
                    **parse_score(cell.value), **trace(sheet, cell)})

        for suffix, target in [("AREA", areas), ("SG", groups)]:
            sheet = qilt[f"FOCUS_{code}_UNI_1Y_{suffix}"]
            used_sheets.add(sheet.title)
            for row in body(sheet):
                for column, metric in enumerate(METRICS, 2):
                    cell = row[column]
                    target.append({"group_type": str(row[0].value or "Study area").strip(),
                        "group": str(row[1].value).strip(), "level": level, "year": 2025,
                        "metric": metric, **parse_score(cell.value), **trace(sheet, cell)})

        sheet = qilt[f"BELONG_{code}_ALL_17-YY_HEPTYPE"]
        used_sheets.add(sheet.title)
        for row in body(sheet):
            for column in range(2, 11):
                cell = row[column]
                belonging.append({"provider_type": row[1].value, "level": level,
                    "year": 2017 + column - 2, **parse_score(cell.value), **trace(sheet, cell)})

        for family in ["SUPP", "TEACH", "RESR", "ENGAG"]:
            sheet = qilt[f"{family}_{code}_UNI_17-YY_STAGE"]
            used_sheets.add(sheet.title)
            for row in body(sheet):
                for cell in row[2:]:
                    header = str(sheet.cell(4, cell.column).value or "")
                    match = re.match(r"(\d{4}):\s*(.+)", header)
                    if not match or int(match[1]) != 2025:
                        continue
                    items.append({"family": family, "item": row[1].value, "level": level,
                        "stage": match[2], "year": 2025, **parse_score(cell.value), **trace(sheet, cell)})

        sheet = qilt[f"CONSID_{code}_UNI_17-YY_CH"]
        used_sheets.add(sheet.title)
        for row in body(sheet):
            for column in range(2, 11):
                cell = row[column]
                leaving.append({"reason": row[1].value, "level": level,
                    "year": 2017 + column - 2, **parse_score(cell.value), **trace(sheet, cell)})

    for name, rows in [("university_scores.csv", universities), ("national_benchmarks.csv", benchmarks),
                       ("study_area_scores.csv", areas), ("student_group_scores.csv", groups),
                       ("belonging_trends.csv", belonging), ("experience_items.csv", items),
                       ("leaving_reasons.csv", leaving)]:
        outputs.append(write_csv(out, name, rows))

    enrolments = []
    sheet = education["2.5"]
    for row in sheet.iter_rows(min_row=4, max_row=53):
        if "Non-University" in str(row[1].value):
            continue
        record = {"institution": clean_name(row[1].value), "institution_state": row[0].value, "year": 2024}
        for column, field in [(2, "research"), (3, "postgraduate_coursework"), (4, "bachelor"),
                              (5, "sub_bachelor"), (6, "enabling"), (7, "non_award"), (8, "total")]:
            record[field], record[field + "_status"] = count(row[column].value)
        record.update({"source_sheet": "2.5", "source_row": row[0].row})
        enrolments.append(record)
    outputs.append(write_csv(out, "institution_enrolments.csv", enrolments))

    flows = []
    sheet = education["2.4"]
    for row in sheet.iter_rows(min_row=5, max_row=14):
        for column in range(1, 10):
            cell = row[column]
            students, status = count(cell.value)
            flows.append({"home_state": row[0].value, "institution_state": sheet.cell(4, column + 1).value,
                "year": 2024, "students": students, "status": status, **trace(sheet, cell)})
    outputs.append(write_csv(out, "home_institution_flows.csv", flows))

    # Use the published row total; suppressed flows and Multi-State providers
    # stay in this denominator. This is an administrative state comparison.
    shares = []
    for row in sheet.iter_rows(min_row=5, max_row=12):
        state = row[0].value
        same_state = row[1 + STATES.index(state)].value
        total = row[10].value
        shares.append({"home_state": state, "year": 2024, "same_state_students": same_state,
            "total_students": total, "same_state_pct": round(100 * same_state / total, 6),
            "source_sheet": "2.4", "same_state_cell": f"{get_column_letter(2 + STATES.index(state))}{row[0].row}",
            "total_cell": f"K{row[0].row}"})
    outputs.append(write_csv(out, "home_state_shares.csv", shares))

    qilt_names = {r["institution"] for r in universities}
    enrolment_keys = {name_key(r["institution"]) for r in enrolments}
    unmatched = sorted(name for name in qilt_names if name_key(name) not in enrolment_keys)
    sheet_metadata = {}
    for title in sorted(used_sheets):
        selected = qilt[title]
        last_body_row = max(row[0].row for row in body(selected))
        sheet_metadata[title] = {"title": selected.cell(1, 1).value,
            "notes": [str(row[0].value) for row in selected.iter_rows(min_row=last_body_row + 1)
                      if row[0].value is not None and str(row[0].value).strip()]}
    manifest = {"verified_date": "2026-10-19", "sources": SOURCES, "metrics": LABELS,
        "percentage_units": "0 to 100, not 0 to 1", "outputs": outputs,
        "qilt_universities": len(qilt_names), "unmatched_enrolment_names": unmatched,
        "sheet_metadata": sheet_metadata}
    (out / "sources.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    assert len(shares) == 8 and len(flows) == 90
    assert all(0 <= r["same_state_pct"] <= 100 for r in shares)
    assert len(qilt_names) == 43
    print(json.dumps({"outputs": outputs, "qilt_universities": len(qilt_names),
                      "unmatched_enrolment_names": unmatched}, indent=2))


if __name__ == "__main__":
    main()
